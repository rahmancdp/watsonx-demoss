import streamlit as st
import subprocess
import orjson

from ibm_watsonx_ai.foundation_models.utils.enums import ModelTypes
from ibm_watsonx_ai.foundation_models import Model
from ibm_watsonx_ai.metanames import GenTextParamsMetaNames as GenParams
from ibm_watsonx_ai.foundation_models.extensions.langchain import WatsonxLLM

import pandas as pd
import tempfile
import pathlib
import json

import gen_java, gen_spec

def split_smaller(text):
    return text

# from modernization_assistant.code_sample.assemblerCode import split_assem_code, clean_assem_code
# from javaCode import split_java_code
# from code_sample.cobolCode import split_smaller
from docx import Document

import os
from dotenv import load_dotenv

load_dotenv()

api_key = os.getenv("API_KEY", None)
project_id = os.getenv("PROJECT_ID", None)

creds = {
    "url"    : "https://us-south.ml.cloud.ibm.com",
    "apikey" : api_key
}

params = {
    GenParams.DECODING_METHOD:"greedy",
    GenParams.MAX_NEW_TOKENS:4000,
    GenParams.MIN_NEW_TOKENS:1,
    GenParams.TOP_K:50,
    GenParams.TOP_P:1,
    GenParams.STOP_SEQUENCES:['```'],
}

paramsmixtral = {
    GenParams.DECODING_METHOD:"greedy",
    GenParams.MAX_NEW_TOKENS: 4000,
    GenParams.MIN_NEW_TOKENS:1,
    GenParams.TOP_K:50,
    GenParams.TOP_P:1,
    GenParams.STOP_SEQUENCES:['```'],
}

import ssl

try:
    _create_unverified_https_context = ssl._create_unverified_context
except AttributeError:
    pass
else:
    ssl._create_default_https_context = _create_unverified_https_context


llmmistra = Model("ibm-mistralai/mixtral-8x7b-instruct-v01-q",creds, paramsmixtral, project_id)
# llmllama = Model("ibm/granite-13b-instruct-v2",creds, params, project_id)
llmllama = Model("meta-llama/llama-2-70b-chat",creds, params, project_id)

def split_string(string, chunk_size):
    return [string[i:i+chunk_size] for i in range(0, len(string), chunk_size)]

def merge_strings(string1, string2):
    prompt = f"""[INST]merge the following strings in context:
    -remove the overlapped
    <<SYS>>
    string1:`{string1}`
    string2:`{string2}`
    <</SYS>>
    [/INST]merged:"""
    answer = ""
    for chunk in llmllama.generate_text([prompt]):
        answer += chunk
    return answer

def merge_string_array(strings):
    # Initialize the merged string with the first element
    if not strings:
        return ""
    merged = strings[0]
    
    # Iterate through the string array and merge each pair
    for next_string in strings[1:]:
        merged = merge_strings(merged, next_string)
    
    return merged

temp_dir = tempfile.TemporaryDirectory()

st.set_page_config(layout="wide")

if "df" not in st.session_state:
    st.session_state.df = None
if "pseudocode" not in st.session_state:
    st.session_state.pseudocode = []
if "summary" not in st.session_state:
    st.session_state.summary = []

if "dfinputs" not in st.session_state:
    st.session_state.dfinputs = None
if "dfoutputs" not in st.session_state:
    st.session_state.dfoutputs = None
if "dfprocesses" not in st.session_state:
    st.session_state.dfprocesses = None

if "funcspec" not in st.session_state:
    st.session_state.funcspec = None
if "inputinquery" not in st.session_state:
    st.session_state.inputinquery = ""
if "outputinquery" not in st.session_state:
    st.session_state.outputinquery = ""
if "processinquery" not in st.session_state:
    st.session_state.processinquery = ""
if "inputmaintenance" not in st.session_state:
    st.session_state.inputmaintenance = ""
if "processmaintenance" not in st.session_state:
    st.session_state.processmaintenance = ""
if "processvalidate" not in st.session_state:
    st.session_state.processvalidate = ""
if "journalconversion" not in st.session_state:
    st.session_state.journalconversion = ""
if "C101" not in st.session_state:
    st.session_state.c101 = ""
if "C102" not in st.session_state:
    st.session_state.c102 = ""
if "version" not in st.session_state:
    st.session_state.version = ""
if "query" not in st.session_state:
    st.session_state.query = ""
if "raml" not in st.session_state:
    st.session_state.raml = ""
if "javaout" not in st.session_state:
    st.session_state.javaout = ""
if "javawithddd" not in st.session_state:
    st.session_state.javawithddd = ""
if "outdoc" not in st.session_state:
    st.session_state.outdoc = ""
if "explaination" not in st.session_state:
    st.session_state.explaination = ""

with st.sidebar:
    st.subheader("ModAssist powered by watsonx")

    uploaded_file = st.file_uploader(label="upload the Pseudo Code",type=['doc','docx'])
    if uploaded_file is not None:
        st.write("filename:", uploaded_file.name)
        with open(os.path.join(pathlib.Path(temp_dir.name),uploaded_file.name),"wb") as f:
            f.write(uploaded_file.getbuffer())
            # st.session_state.pseudocode = uploaded_file.read().decode("utf-8", errors="ignore")
            doc = Document(uploaded_file)
            page_content = []
            current_page = []

            for paragraph in doc.paragraphs:
                if paragraph.text:  # Skip empty paragraphs
                    current_page.append(paragraph.text)
                elif current_page:  # Store the current page if it's not empty
                    page_content.append("\n".join(current_page))
                    print("\n".join(current_page))
                    print("#####################")
                    current_page = []

            # Add the last page content if it exists
            if current_page:
                page_content.append("\n".join(current_page))
                print("\n".join(current_page))
                print("#####################")

        print(f"total page {len(page_content)}")
        st.session_state.pseudocode = page_content

    uploaded_file = st.file_uploader(label="upload the Functional Spec",type=['xls','xlsx'])
    if uploaded_file is not None:
        st.write("filename:", uploaded_file.name)
        with open(os.path.join(pathlib.Path(temp_dir.name),uploaded_file.name),"wb") as f:
            f.write(uploaded_file.getbuffer())
            st.session_state.funcspec = pd.read_excel(uploaded_file, sheet_name=None)

tabPseudoCodeToFuncSpec, tabFuncSpecToRAML, tabRAMLToJava = st.tabs(['pseudocode to funcspec','funcspec to RAML/Process','RAML to java'])

with tabPseudoCodeToFuncSpec:
    st.subheader(f"Pseudo Code to Functional Specification")

    # st.session_state.pseudocode = pseudocode_to_funcspec.read_docx('pseudocode/C101 As-Is readable v0.9 - Share to IBM.docx')
    funcspec = ""

    btGenFuncSpecSummary = st.button("Generate Functional Specification Summary")
    btGenFuncSpecXlS = st.button("Generate Functional Specification XLS")

    colpseudocode, colfuncspec = st.columns(2)
    
    with colpseudocode:
        dfpseudocode = pd.DataFrame(st.session_state.pseudocode)
        dfpseudocode = dfpseudocode.applymap(lambda x: x.replace('\n', '<br>'))

        dffuncspec = pd.DataFrame(st.session_state.summary)
        dffuncspec = dffuncspec.applymap(lambda x: x.replace('\n', '<br>'))

        dfall = pd.concat([dfpseudocode,dffuncspec], axis=1)

        st.markdown(dfall.to_html(escape=False), unsafe_allow_html=True)
        # st.dataframe(df)
        # st.markdown(st.session_state.pseudocode)
 
    with colfuncspec:
        st.dataframe(st.session_state.df)

        if btGenFuncSpecSummary:

            st.session_state.summary = []
            for code in st.session_state.pseudocode:
                if code.strip() == '':
                    continue
                chunks = split_string(code,3000)
                funcspecparts = gen_spec.generate_funcspec(llmmistra,chunks)
                
                funcspec = merge_string_array(funcspecparts)

                st.session_state.summary.append(funcspec)
                st.markdown(funcspec)
                st.text("##########")
    
        if btGenFuncSpecXlS:
            index = 0

            for code in st.session_state.pseudocode:
                index += 1
                if index > 4:
                    break
                if code.strip() == '':
                    continue
                chunks = split_string(code,3000)
                funcspecparts = gen_spec.generate_funcspec_json(llmmistra,chunks)

                funcspec = merge_string_array(funcspecparts)

                funcspecjson = funcspec.replace("```","")
                print("$$$$"+funcspecjson)
                
                try:
                    data_dict = orjson.loads(funcspecjson)
                    # Create a DataFrame from the dictionary

                    dfinputs = pd.DataFrame(data_dict["inputs"])
                    if st.session_state.dfinputs is None:
                        st.session_state.dfinputs = dfinputs
                    else:
                        print("appending input")
                        st.session_state.dfinputs = st.session_state.dfinputs.append(dfinputs)
                        # st.session_state.dfinputs = pd.merge(st.session_state.dfinputs, dfinputs, on='Field Description')
                        
                    # st.session_state.dfinputs = st.session_state.dfinputs.reset_index(drop=True)

                    st.dataframe(st.session_state.dfinputs)

                    dfoutputs = pd.DataFrame(data_dict["outputs"])
                    if st.session_state.dfoutputs is None:
                        st.session_state.dfoutputs = dfoutputs
                    else:
                        print("appending output")
                        st.session_state.dfoutputs = st.session_state.dfoutputs.append(dfoutputs)
                        # st.session_state.dfoutputs = pd.merge(st.session_state.dfoutputs, dfoutputs, on='Field Description')
                        
                    st.dataframe(st.session_state.dfoutputs)

                    dfprocesses = pd.DataFrame(data_dict["processes"])
                    if st.session_state.dfprocesses is None:
                        st.session_state.dfprocesses = dfprocesses
                    else:
                        print("appending process")
                        st.session_state.dfprocesses = st.session_state.dfprocesses.append(dfprocesses)
                        # st.session_state.dfprocesses = st.session_state.dfprocesses.reset_index(drop=True)

                    st.dataframe(st.session_state.dfprocesses)
                except json.JSONDecodeError as e:
                    st.text(f"JSON decoding error: {str(e)}")
                    st.text('problem in json: '+funcspecjson)
                st.text("##########")

            gen_spec.save_to_excel_multiple('output/test.xlsx',st.session_state.dfinputs,st.session_state.dfoutputs,st.session_state.dfprocesses)
            button_text = 'Download Functional Specification File'
            with open(output_file_path,'rb') as xls:
                st.download_button(button_text, xls, 'funcspec.xlsx')
        # st.dataframe(st.session_state.dfinputs)
        # st.dataframe(st.session_state.dfoutputs)
        # st.dataframe(st.session_state.dfprocesses)


with tabFuncSpecToRAML:
    st.subheader(f"Functional Specification to RAML")


    # st.session_state.funcspec = pd.read_excel('funcspec/Functional Spec - Account Statement & Advice control service - Share to IBM.xlsx', sheet_name=None)
    if st.session_state.funcspec is not None:
        sheetnames = []
        for sheet_name, df in st.session_state.funcspec.items():
            sheetnames.append(sheet_name)

        tabs = st.tabs(sheetnames)
        tabindex = 0
        for sheet_name, df in st.session_state.funcspec.items():
            with tabs[tabindex]:
                st.text(f"Sheet Name: {sheet_name}")
                st.dataframe(df)

                if st.button(f"Generate {sheet_name} RAML/Java"):
                    if sheet_name.startswith("Input(Inq)"):
                        st.session_state.inputinquery = gen_java.generate_inputfunction_raml(llmllama,df)
                        st.code(st.session_state.inputinquery,language="yaml")
                        st.text("##########")
                    elif sheet_name.startswith("Output(Inq)"):
                        st.session_state.outputinquery = gen_java.generate_outputfunction_raml(llmllama,df)
                        st.code(st.session_state.outputinquery,language="yaml")
                        st.text("##########")
                    elif sheet_name.startswith("Process(Inq)"):
                        st.session_state.processinquery = gen_java.generate_processinquery(llmllama,df)
                        st.code(st.session_state.processinquery,language="java")
                        st.text("##########")
                    elif sheet_name.startswith("Input(Maint)"):
                        st.session_state.inputmaintenance = gen_java.generate_inputfunction_raml(llmllama,df)
                        st.code(st.session_state.inputmaintenance,language="yaml")
                        st.text("##########")
                    elif sheet_name.startswith("Process(Maint)"):
                        st.session_state.processmaintenance = gen_java.generate_processmaintenance(llmllama,df)
                        st.code(st.session_state.processmaintenance,language="java")
                        st.text("##########")
                    elif sheet_name.startswith("Process(Validate)"):
                        st.session_state.processvalidate = gen_java.generate_processmaintenance(llmllama,df)
                        st.code(st.session_state.processvalidate,language="java")
                        st.text("##########")
                    elif sheet_name.startswith("Journal Conversion"):
                        screenfunction = gen_java.generate_screenfunction_raml(llmllama,df)
                        st.code(screenfunction,language="yaml")
                        screenfunction = generate_screenfunction(df)
                        st.code(screenfunction,language="java")
                        st.text("##########")
                    elif sheet_name.startswith("C101"):
                        screenfunction = gen_java.generate_screenfunction_raml(llmllama,df)
                        st.code(screenfunction,language="yaml")
                        screenfunction = generate_screenfunction(df)
                        st.code(screenfunction,language="java")
                        st.text("##########")
                    elif sheet_name.startswith("C102"):
                        screenfunction = gen_java.generate_screenfunction_raml(llmllama,df)
                        st.code(screenfunction,language="yaml")
                        screenfunction = generate_screenfunction(df)
                        st.code(screenfunction,language="java")
                        st.text("##########")
                    elif sheet_name.startswith("version"):
                        screenfunction = gen_java.generate_screenfunction_raml(llmllama,df)
                        st.code(screenfunction,language="yaml")
                        screenfunction = generate_screenfunction(df)
                        st.code(screenfunction,language="java")
                        st.text("##########")
                    elif sheet_name.startswith("query"):
                        screenfunction = gen_java.generate_screenfunction_raml(llmllama,df)
                        st.code(screenfunction,language="yaml")
                        screenfunction = generate_screenfunction(df)
                        st.code(screenfunction,language="java")
                        st.text("##########")

                tabindex += 1

with tabRAMLToJava:
    st.subheader(f"Generate Java Skeleton")
    if st.button("Generate Swagger"):
        for raml in [st.session_state.inputinquery,st.session_state.outputinquery,st.session_state.inputmaintenance]:
            st.code(raml,language="yaml")
            swagger = gen_java.generate_swagger(llmllama,raml)
            st.code(swagger,language="json")
            
            jsonfile = swagger.encode('utf-8')

            st.download_button(
                label="Download JSON",
                data=jsonfile,
                file_name='swagger.json',
                mime='text/json',
            )

            with open("swagger.json", "w") as file:
                file.write(swagger)

            # Execute the command with the saved Swagger JSON file
            command = "swagger-codegen generate -i swagger.json \
-l spring --api-package com.hsbc.hbap.obs.transaction.application.controller.command \
--model-package com.hsbc.hbap.obs.transaction.application.vo \
--invoker-package=com.hsbc.hbap.obs.transaction \
--additional-properties configPackage=com.hsbc.hbap.obs.transaction.config useTags=true -o out"
            print(command)
            result = subprocess.run(command, shell=True, capture_output=True, text=True)
            
            # Display the command output
            st.code(result.stdout)